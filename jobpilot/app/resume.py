"""Resume handling. Your resume is stored in the visible  jobpilot/resume/  folder.
Drop a file there yourself or upload it in the dashboard - both work. Older resumes move to resume/previous/."""
import re
import shutil
from pathlib import Path

from . import config, llm
from .filters import find_city

SKILLS = [
    # cloud / devops
    "aws", "azure", "gcp", "docker", "kubernetes", "terraform", "ansible", "jenkins", "git", "github actions", "gitlab ci",
    "ci/cd", "linux", "bash", "shell scripting", "nginx", "prometheus", "grafana", "cloudwatch", "ec2", "s3", "vpc", "iam",
    "lambda", "eks", "ecs", "cloudformation", "helm", "argocd", "devops", "sre", "networking",
    # programming / web
    "python", "java", "javascript", "typescript", "node.js", "react", "angular", "vue", "html", "css", "php", "c++", "c#",
    ".net", "spring boot", "django", "flask", "fastapi", "express", "rest api", "microservices", "golang", "kotlin", "swift",
    "android", "flutter", "react native", "wordpress",
    # data / ai
    "sql", "mysql", "postgresql", "mongodb", "oracle", "redis", "pandas", "numpy", "machine learning", "deep learning",
    "data analysis", "data science", "power bi", "tableau", "excel", "etl", "spark", "hadoop", "nlp", "tensorflow", "pytorch",
    # testing / other
    "selenium", "manual testing", "automation testing", "api testing", "postman", "jira", "agile", "scrum", "figma",
    "ui/ux", "seo", "digital marketing", "tally", "sap", "salesforce", "communication", "customer support", "content writing",
    "accounting", "gst", "recruitment", "business development", "sales",
]

ROLE_MAP = [   # (skills that signal it, roles to search)
    ({"aws", "azure", "gcp", "docker", "kubernetes", "terraform", "jenkins", "devops", "ci/cd"},
     ["DevOps Engineer", "Cloud Engineer", "Junior DevOps Engineer", "Site Reliability Engineer", "AWS Engineer"]),
    ({"react", "angular", "vue", "html", "css", "javascript", "typescript"},
     ["Frontend Developer", "Web Developer", "React Developer", "Full Stack Developer"]),
    ({"node.js", "express", "django", "flask", "fastapi", "spring boot", "java", "python", "php", ".net"},
     ["Software Engineer", "Backend Developer", "Python Developer", "Java Developer"]),
    ({"machine learning", "deep learning", "data science", "tensorflow", "pytorch", "nlp"},
     ["Data Scientist", "Machine Learning Engineer", "AI Engineer"]),
    ({"data analysis", "power bi", "tableau", "excel", "sql", "pandas"},
     ["Data Analyst", "Business Analyst", "Power BI Developer"]),
    ({"selenium", "manual testing", "automation testing", "api testing"},
     ["QA Engineer", "Software Tester", "Automation Test Engineer"]),
    ({"android", "flutter", "react native", "kotlin", "swift"}, ["Mobile App Developer", "Flutter Developer", "Android Developer"]),
    ({"seo", "digital marketing", "content writing"}, ["Digital Marketing Executive", "SEO Executive", "Content Writer"]),
    ({"tally", "accounting", "gst"}, ["Accountant", "Accounts Executive"]),
    ({"recruitment"}, ["HR Recruiter", "Talent Acquisition Executive"]),
    ({"customer support", "communication", "sales", "business development"}, ["Customer Support Executive", "Business Development Executive"]),
]


# ---------------------------------------------------------------- files
def _safe_name(name):
    stem = re.sub(r"[^\w.\- ]+", "_", Path(name).stem).strip(" ._") or "resume"
    return stem[:80] + Path(name).suffix.lower()


def list_files():
    return sorted([p for p in config.RESUME_DIR.iterdir() if p.is_file() and p.suffix.lower() in config.RESUME_EXTS],
                  key=lambda p: p.stat().st_mtime, reverse=True)


def active_path(saved=None):
    """The resume used for applying: the saved one if it still exists, else the newest file in the folder."""
    if saved and Path(saved).exists():
        return Path(saved)
    files = list_files()
    return files[0] if files else None


def save_upload(fileobj, filename):
    ext = Path(filename).suffix.lower()
    if ext not in config.RESUME_EXTS:
        raise ValueError("Upload a PDF, DOCX or TXT resume")
    prev = config.RESUME_DIR / "previous"
    for old in list_files():
        prev.mkdir(exist_ok=True)
        shutil.move(str(old), str(prev / old.name))
    dest = config.RESUME_DIR / _safe_name(filename)
    with dest.open("wb") as f:
        shutil.copyfileobj(fileobj, f)
    return dest


def remove_active():
    prev = config.RESUME_DIR / "previous"
    for old in list_files():
        prev.mkdir(exist_ok=True)
        shutil.move(str(old), str(prev / old.name))


# ---------------------------------------------------------------- reading
def extract_text(path):
    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        return "\n".join((pg.extract_text() or "") for pg in PdfReader(str(p)).pages)
    if ext == ".docx":
        import docx

        d = docx.Document(str(p))
        parts = [para.text for para in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    return p.read_text(errors="ignore")


def _years(text):
    best = 0.0
    for m in re.finditer(r"(\d{1,2}(?:\.\d)?)\s*\+?\s*(?:years?|yrs?)\b", text, re.I):
        v = float(m.group(1))
        if v <= 40:
            best = max(best, v)
    return best


def _regex_profile(text):
    email = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)
    phone = re.search(r"(?:\+?91[\s-]?)?[6-9]\d{4}[\s-]?\d{5}", text) or re.search(r"\+?\d[\d\s().-]{8,15}\d", text)
    linkedin = re.search(r"(?:https?://)?(?:www\.)?linkedin\.com/in/[\w-]+/?", text, re.I)
    github = re.search(r"(?:https?://)?(?:www\.)?github\.com/[\w-]+", text, re.I)
    low = text.lower()
    skills = [s for s in SKILLS if re.search(r"(?<![\w])" + re.escape(s) + r"(?![\w+#])", low)]
    name = ""
    for ln in text.splitlines()[:8]:
        ln = ln.strip()
        if ln and "@" not in ln and not re.search(r"\d|resume|curriculum|vitae|\bcv\b", ln, re.I) and len(ln.split()) <= 4:
            name = ln.title() if ln.isupper() else ln
            break
    parts = name.split()
    ph = re.sub(r"[^\d+]", "", phone.group(0)) if phone else ""
    if len(ph) == 12 and ph.startswith("91"):
        ph = ph[2:]
    return {
        "first_name": parts[0] if parts else "", "last_name": " ".join(parts[1:]) if len(parts) > 1 else "",
        "email": email.group(0) if email else "", "phone": ph[-10:] if len(ph) >= 10 else ph,
        "linkedin": linkedin.group(0) if linkedin else "", "github": github.group(0) if github else "",
        "website": "", "location": find_city(text[:1500]) or "", "current_company": "", "current_title": "",
        "years_experience": _years(text), "skills": skills, "target_roles": [], "summary": "", "education": "",
        "notice_period": "", "expected_ctc": "", "current_ctc": "",
    }


PROMPT = """Extract a job-seeker profile from this resume (Indian job market). Reply with ONLY one JSON object:
{"first_name":"","last_name":"","email":"","phone":"10 digit mobile","linkedin":"","github":"","website":"",
 "location":"city, India","current_company":"","current_title":"","years_experience":0,
 "skills":["every concrete technical/professional skill, tool, cert as short lowercase keywords"],
 "target_roles":["6-8 job titles this person should apply for in India, most fitting first, include fresher/junior variants if early career"],
 "summary":"2 sentences","education":"highest degree, field, year",
 "notice_period":"e.g. 30 days / Immediate / 15 days if written, else empty",
 "expected_ctc":"e.g. 12 LPA if written, else empty",
 "current_ctc":"e.g. 8 LPA if written, else empty"}
Use "" or 0 when something is not in the resume. Never invent facts. Prefer Indian formats (LPA, notice in days).

RESUME:
"""


def parse_profile(text):
    base = _regex_profile(text)
    data = llm.complete_json("You are a precise resume parser.", PROMPT + text[:12000], 2000)
    if data:
        for k, v in data.items():
            if k in base and v not in (None, "", [], 0):
                base[k] = v
    if not base["target_roles"]:
        sk = set(base["skills"])
        roles = []
        for sig, rs in ROLE_MAP:
            if sk & sig:
                roles += rs[:3]
        if base["current_title"]:
            roles.insert(0, base["current_title"])
        base["target_roles"] = list(dict.fromkeys(roles))[:8]
    base["skills"] = sorted({s.strip().lower() for s in base["skills"] if s.strip()})
    try:
        base["years_experience"] = float(base["years_experience"] or 0)
    except (TypeError, ValueError):
        base["years_experience"] = 0
    return base
